"""
Assemble paper-style figure plates into a PowerPoint deck.

Composes multi-panel plates (panel letters A/B/C…) with PIL, then places one
plate per slide with python-pptx. Cohort I = main figures, Cohort II =
replication. Redundant figures (old combined fig1D / fig6, and fig1D which
overlaps fig4A) are excluded.

Usage:
    python python/make_pptx.py
"""

import os
from PIL import Image, ImageDraw, ImageFont

Image.MAX_IMAGE_PIXELS = None   # plates are our own composites, not untrusted input
MAX_PLATE_PX = 4000             # cap longest plate edge before saving

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

FIG_DIR = os.path.join(os.path.dirname(__file__), "..", "results", "figures_solely")
FIG_DIR = os.path.abspath(FIG_DIR)
OUT_PPTX = os.path.join(FIG_DIR, "BioMe_AM_figures.pptx")

PAD = 30          # padding between panels (px)
LETTER_PT = 64    # panel letter size
BG = (255, 255, 255)


def _font(size, bold=True):
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/Library/Fonts/Arial.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
    ]
    for c in candidates:
        if os.path.exists(c):
            try:
                return ImageFont.truetype(c, size)
            except Exception:
                continue
    return ImageFont.load_default()


def _load(name):
    p = os.path.join(FIG_DIR, name)
    if not os.path.exists(p):
        raise FileNotFoundError(p)
    return Image.open(p).convert("RGB")


def _label(img, letter):
    """Draw a bold panel letter in the top-left corner."""
    img = img.copy()
    draw = ImageDraw.Draw(img)
    font = _font(LETTER_PT)
    draw.text((12, 6), letter, fill=(0, 0, 0), font=font)
    return img


def _resize_h(img, h):
    w = int(img.width * h / img.height)
    return img.resize((w, h), Image.LANCZOS)


def _resize_w(img, w):
    h = int(img.height * w / img.width)
    return img.resize((w, h), Image.LANCZOS)


def hstack(imgs):
    h = max(i.height for i in imgs)
    imgs = [_resize_h(i, h) for i in imgs]
    w = sum(i.width for i in imgs) + PAD * (len(imgs) - 1)
    plate = Image.new("RGB", (w, h), BG)
    x = 0
    for i in imgs:
        plate.paste(i, (x, 0))
        x += i.width + PAD
    return plate


def vstack(imgs):
    w = max(i.width for i in imgs)
    imgs = [_resize_w(i, w) for i in imgs]
    h = sum(i.height for i in imgs) + PAD * (len(imgs) - 1)
    plate = Image.new("RGB", (w, h), BG)
    y = 0
    for i in imgs:
        plate.paste(i, (0, y))
        y += i.height + PAD
    return plate


def build_plates(cohort):
    """Return list of (figure_title, PIL plate) for one cohort."""
    c = cohort
    plates = []

    # Figure 1 — Overview: A venn, B variant counts, C carrier frequency
    fig1 = hstack([
        _label(_load(f"fig1A_venn_{c}.png"), "A"),
        _label(_load(f"fig1B_variant_counts_{c}.png"), "B"),
        _label(_load(f"fig1C_carrier_freq_{c}.png"), "C"),
    ])
    plates.append(("Figure 1. Variant and carrier overview", fig1))

    # Figure 2 — Carrier frequency by gene (single wide panel)
    plates.append(("Figure 2. Carrier frequency by gene",
                   _load(f"fig2_carrier_by_gene_{c}.png")))

    # Figure 3 — Regression table (one slide per page)
    if c == "cohortI":
        for pg in (1, 2, 3):
            f = f"fig3_regression_table_p{pg}_{c}.png"
            if os.path.exists(os.path.join(FIG_DIR, f)):
                plates.append((f"Figure 3. Logistic regression results ({pg}/3)",
                               _load(f)))
    else:
        plates.append(("Figure 3. Logistic regression results",
                       _load(f"fig3_regression_table_{c}.png")))

    # Figure 4 — A phenotype barplot (own slide); B gene x phenotype heatmaps
    plates.append(("Figure 4A. Carrier frequency by phenotype group",
                   _load(f"fig4A_barplot_{c}.png")))
    fig4_b = hstack([
        _label(_load(f"fig4B_heatmap_ACMGplp_{c}.png"), "i"),
        _label(_load(f"fig4B_heatmap_AMcalibrated_{c}.png"), "ii"),
        _label(_load(f"fig4B_heatmap_AMcalibratedNotPLP_{c}.png"), "iii"),
    ])
    plates.append(("Figure 4B. Gene × phenotype carrier frequency", fig4_b))

    # Figure 5 — A ancestry bar (own slide); B ancestry heatmaps
    plates.append(("Figure 5A. Carrier frequency by ancestry",
                   _load(f"fig5A_ancestry_bar_{c}.png")))
    fig5_b = hstack([
        _label(_load(f"fig5B_heatmap_ACMGplp_{c}.png"), "i"),
        _label(_load(f"fig5B_heatmap_AMcalibrated_{c}.png"), "ii"),
        _label(_load(f"fig5B_heatmap_AMcalibratedNotPLP_{c}.png"), "iii"),
    ])
    plates.append(("Figure 5B. Gene × ancestry carrier frequency", fig5_b))

    # Figure 6 — Forest plots (3 variant classes)
    fig6 = hstack([
        _label(_load(f"fig6_forest_ACMGplp_{c}.png"), "A"),
        _label(_load(f"fig6_forest_AMcalibrated_{c}.png"), "B"),
        _label(_load(f"fig6_forest_AMcalibratedNotPLP_{c}.png"), "C"),
    ])
    plates.append(("Figure 6. Phenotype associations (OR, 95% CI)", fig6))

    return plates


def add_slide(prs, title, plate_path):
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)

    # Title
    tb = slide.shapes.add_textbox(Inches(0.4), Inches(0.15),
                                  prs.slide_width - Inches(0.8), Inches(0.6))
    tf = tb.text_frame
    tf.text = title
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # Image area below title
    img = Image.open(plate_path)
    avail_w = prs.slide_width - Inches(0.8)
    avail_h = prs.slide_height - Inches(1.1)
    scale = min(avail_w / img.width, avail_h / img.height)
    w = int(img.width * scale)
    h = int(img.height * scale)
    left = int((prs.slide_width - w) / 2)
    top = Inches(0.95)
    slide.shapes.add_picture(plate_path, left, top, width=w, height=h)


def title_slide(prs):
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    tb = slide.shapes.add_textbox(Inches(1), Inches(2.6),
                                  prs.slide_width - Inches(2), Inches(2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.text = "AlphaMissense Calibration in BioMe Hereditary Cancer Genes"
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.runs[0]
    r.font.size = Pt(34)
    r.font.bold = True
    r.font.name = "Arial"
    r.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = "Cohort I (N=28,310) · Cohort II (N=13,967)"
    r2.font.size = Pt(18)
    r2.font.name = "Arial"
    r2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def section_slide(prs, text):
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    tb = slide.shapes.add_textbox(Inches(1), Inches(3.1),
                                  prs.slide_width - Inches(2), Inches(1.3))
    tf = tb.text_frame
    tf.text = text
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.runs[0]
    r.font.size = Pt(30)
    r.font.bold = True
    r.font.name = "Arial"
    r.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    tmp_dir = os.path.join(FIG_DIR, "_plates_tmp")
    os.makedirs(tmp_dir, exist_ok=True)

    title_slide(prs)

    for cohort, sect in [("cohortI", "Cohort I — Discovery"),
                         ("cohortII", "Cohort II — Replication")]:
        section_slide(prs, sect)
        for i, (title, plate) in enumerate(build_plates(cohort)):
            # Downscale so the longest edge is <= MAX_PLATE_PX
            longest = max(plate.width, plate.height)
            if longest > MAX_PLATE_PX:
                s = MAX_PLATE_PX / longest
                plate = plate.resize((int(plate.width * s), int(plate.height * s)),
                                     Image.LANCZOS)
            plate_path = os.path.join(tmp_dir, f"{cohort}_{i:02d}.png")
            plate.save(plate_path, dpi=(200, 200))
            add_slide(prs, f"{title} — {sect.split(' — ')[0]}", plate_path)

    prs.save(OUT_PPTX)
    print(f"Saved: {OUT_PPTX}")
    print(f"Slides: {len(prs.slides._sldIdLst)}")


if __name__ == "__main__":
    main()
